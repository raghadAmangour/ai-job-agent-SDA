"""Resume text extraction — identical logic to Phase 2 notebook.
Supports PDF (text-based, not scanned), DOCX, and TXT/MD files."""
from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument

MIN_RESUME_CHARS = 200


def read_pdf(path) -> str:
    parts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            parts.append(text)
    return "\n".join(parts)


def read_docx(path) -> str:
    doc = DocxDocument(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_txt(path) -> str:
    return Path(path).read_text(encoding="utf-8", errors="ignore")


def read_resume(path) -> str:
    """Dispatch by extension. Returns raw extracted text (not yet cleaned)."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix in (".txt", ".md"):
        return read_txt(path)
    raise ValueError(
        f"Unsupported file type: {suffix}. Use PDF, DOCX, or TXT. "
        f"(.doc or scanned/image PDFs are not supported — convert first)"
    )


def resume_quality_check(raw_text: str) -> dict:
    """Flags likely-scanned or empty PDFs before wasting an LLM call."""
    chars = len(raw_text.strip())
    words = len(raw_text.split())
    return {
        "chars": chars,
        "words": words,
        "likely_scanned_or_empty": chars < MIN_RESUME_CHARS,
    }
